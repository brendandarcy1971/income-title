from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-08.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)
ACCENT = RGBColor(198, 94, 31)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_oval(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_metric_chip(slide, left, top, width, label, value):
    add_rect(slide, left, top, width, 0.34, RGBColor(255, 255, 255), RGBColor(206, 213, 221))
    add_text(slide, left + 0.08, top + 0.05, width - 0.16, 0.12, label, size=7, color=SUB)
    add_text(slide, left + 0.08, top + 0.18, width - 0.16, 0.12, value, size=9, bold=True, color=ACCENT)


def add_callout(slide, left, top, text):
    chip = add_rect(slide, left, top, 2.0, 0.32, RGBColor(255, 255, 255), BORDER)
    _ = chip
    add_text(slide, left + 0.08, top + 0.05, 1.85, 0.22, text, size=9, bold=True, color=TEAL)


def draw_floorplan_panel(slide):
    left = 0.45
    top = 1.35
    width = 7.75
    height = 5.05

    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_text(slide, left + 0.18, top + 0.08, 5.4, 0.28, "Illustrative Marketing Floorplan + Site", size=11, bold=True, color=SUB)

    # Display-home style metrics strip
    metrics_top = top + 0.08
    add_metric_chip(slide, left + 4.86, metrics_top, 0.9, "LOT", "~375 sqm")
    add_metric_chip(slide, left + 5.83, metrics_top, 0.9, "INTERNAL", "~148 sqm")
    add_metric_chip(slide, left + 6.8, metrics_top, 0.9, "PARKING", "1 + drive")

    # Site frame
    plan_left = left + 0.25
    plan_top = top + 0.45
    plan_w = 6.45
    plan_h = 4.0
    add_rect(slide, plan_left, plan_top, plan_w, plan_h, RGBColor(246, 252, 246), RGBColor(170, 182, 170))

    # Marketing-style dimension guides
    add_rect(slide, plan_left, plan_top - 0.08, plan_w, 0.01, RGBColor(176, 184, 194), RGBColor(176, 184, 194))
    add_rect(slide, plan_left, plan_top - 0.13, 0.01, 0.1, RGBColor(176, 184, 194), RGBColor(176, 184, 194))
    add_rect(slide, plan_left + plan_w - 0.01, plan_top - 0.13, 0.01, 0.1, RGBColor(176, 184, 194), RGBColor(176, 184, 194))
    add_text(slide, plan_left + 2.35, plan_top - 0.25, 1.8, 0.13, "Lot width ~12.5m", size=7, color=SUB, align=PP_ALIGN.CENTER)

    add_rect(slide, plan_left - 0.08, plan_top, 0.01, plan_h, RGBColor(176, 184, 194), RGBColor(176, 184, 194))
    add_rect(slide, plan_left - 0.13, plan_top, 0.1, 0.01, RGBColor(176, 184, 194), RGBColor(176, 184, 194))
    add_rect(slide, plan_left - 0.13, plan_top + plan_h - 0.01, 0.1, 0.01, RGBColor(176, 184, 194), RGBColor(176, 184, 194))
    add_text(slide, plan_left - 0.5, plan_top + 1.86, 0.35, 0.2, "Lot depth\n~30m", size=7, color=SUB, align=PP_ALIGN.CENTER)

    # Street edge, driveway, and parking
    street_top = plan_top + plan_h - 0.52
    add_rect(slide, plan_left, street_top, plan_w, 0.52, RGBColor(231, 236, 240), RGBColor(195, 202, 209))
    add_text(slide, plan_left + 4.95, street_top + 0.16, 1.35, 0.2, "Street frontage", size=7, color=SUB)

    add_rect(slide, plan_left + 0.22, street_top - 1.08, 0.96, 1.08, RGBColor(236, 239, 242), RGBColor(198, 204, 212))
    add_text(slide, plan_left + 0.29, street_top - 0.16, 0.8, 0.15, "Driveway", size=7, color=SUB)

    add_rect(slide, plan_left + 0.22, street_top - 1.94, 1.48, 0.86, RGBColor(242, 245, 248), RGBColor(194, 201, 209))
    add_text(slide, plan_left + 0.45, street_top - 1.55, 1.0, 0.18, "Garage", size=8, bold=True, color=SUB)

    # Backyard zone
    add_rect(slide, plan_left + 1.75, plan_top + 0.16, 4.5, 0.86, RGBColor(229, 246, 232), RGBColor(174, 204, 180))
    add_text(slide, plan_left + 3.2, plan_top + 0.48, 1.7, 0.2, "Private Backyard", size=9, bold=True, color=SUB, align=PP_ALIGN.CENTER)

    # Subtle landscaping pattern for a brochure look
    shrub_fill = RGBColor(201, 229, 198)
    shrub_line = RGBColor(161, 197, 160)
    add_oval(slide, plan_left + 1.95, plan_top + 0.3, 0.12, 0.12, shrub_fill, shrub_line)
    add_oval(slide, plan_left + 2.12, plan_top + 0.45, 0.1, 0.1, shrub_fill, shrub_line)
    add_oval(slide, plan_left + 2.28, plan_top + 0.3, 0.12, 0.12, shrub_fill, shrub_line)
    add_oval(slide, plan_left + 5.62, plan_top + 0.3, 0.12, 0.12, shrub_fill, shrub_line)
    add_oval(slide, plan_left + 5.8, plan_top + 0.45, 0.1, 0.1, shrub_fill, shrub_line)
    add_oval(slide, plan_left + 5.95, plan_top + 0.3, 0.12, 0.12, shrub_fill, shrub_line)

    # House footprint
    house_left = plan_left + 1.75
    house_top = plan_top + 1.05
    house_w = 4.5
    house_h = 2.08
    add_rect(slide, house_left, house_top, house_w, house_h, RGBColor(253, 254, 255), RGBColor(169, 178, 188))

    # Internal rooms
    add_rect(slide, house_left + 0.1, house_top + 0.1, 1.28, 0.78, RGBColor(245, 249, 252), RGBColor(184, 194, 204))
    add_text(slide, house_left + 0.22, house_top + 0.42, 1.0, 0.17, "Bed 1", size=8, color=SUB)

    add_rect(slide, house_left + 1.45, house_top + 0.1, 1.1, 0.78, RGBColor(245, 249, 252), RGBColor(184, 194, 204))
    add_text(slide, house_left + 1.72, house_top + 0.42, 0.7, 0.17, "Bed 2", size=8, color=SUB)

    add_rect(slide, house_left + 2.63, house_top + 0.1, 1.02, 0.78, RGBColor(245, 249, 252), RGBColor(184, 194, 204))
    add_text(slide, house_left + 2.86, house_top + 0.42, 0.7, 0.17, "Bed 3", size=8, color=SUB)

    add_rect(slide, house_left + 3.72, house_top + 0.1, 0.68, 0.78, RGBColor(250, 247, 242), RGBColor(195, 198, 205))
    add_text(slide, house_left + 3.83, house_top + 0.42, 0.48, 0.17, "Bath", size=7, color=SUB)

    add_rect(slide, house_left + 0.1, house_top + 0.95, 0.82, 1.03, RGBColor(250, 247, 242), RGBColor(195, 198, 205))
    add_text(slide, house_left + 0.17, house_top + 1.37, 0.65, 0.17, "Ensuite", size=7, color=SUB)

    add_rect(slide, house_left + 0.98, house_top + 0.95, 1.32, 1.03, RGBColor(243, 249, 252), RGBColor(184, 194, 204))
    add_text(slide, house_left + 1.15, house_top + 1.37, 1.0, 0.17, "Living 2", size=8, color=SUB)

    add_rect(slide, house_left + 2.36, house_top + 0.95, 2.04, 1.03, RGBColor(240, 249, 242), RGBColor(178, 197, 182))
    add_text(slide, house_left + 2.66, house_top + 1.23, 1.4, 0.17, "Kitchen + Dining", size=8, color=SUB)
    add_text(slide, house_left + 2.92, house_top + 1.47, 1.0, 0.17, "Family", size=8, bold=True, color=SUB)

    # Patio connection to backyard
    add_rect(slide, house_left + 3.25, house_top - 0.15, 1.08, 0.2, RGBColor(236, 241, 245), RGBColor(198, 206, 213))
    add_text(slide, house_left + 3.45, house_top - 0.13, 0.7, 0.12, "Alfresco", size=7, color=SUB)

    # Mini legend badges
    legend_top = plan_top + 3.5
    add_rect(slide, plan_left + 0.15, legend_top, 0.16, 0.12, RGBColor(245, 249, 252), RGBColor(184, 194, 204))
    add_text(slide, plan_left + 0.35, legend_top - 0.01, 1.05, 0.14, "Internal", size=7, color=SUB)
    add_rect(slide, plan_left + 1.42, legend_top, 0.16, 0.12, RGBColor(229, 246, 232), RGBColor(174, 204, 180))
    add_text(slide, plan_left + 1.62, legend_top - 0.01, 1.05, 0.14, "Outdoor", size=7, color=SUB)
    add_rect(slide, plan_left + 2.62, legend_top, 0.16, 0.12, RGBColor(236, 239, 242), RGBColor(198, 204, 212))
    add_text(slide, plan_left + 2.82, legend_top - 0.01, 1.45, 0.14, "Vehicle access", size=7, color=SUB)

    # Callout chips
    add_callout(slide, left + 0.2, top + 4.55, "3 Bedrooms")
    add_callout(slide, left + 2.28, top + 4.55, "2 Bathrooms")
    add_callout(slide, left + 4.36, top + 4.55, "Single garage")

    add_callout(slide, left + 0.2, top + 4.9, "Private backyard")
    add_callout(slide, left + 2.28, top + 4.9, "Driveway parking")
    add_callout(slide, left + 4.36, top + 4.9, "Indoor-outdoor living")


def draw_spec_panel(slide):
    left = 8.35
    top = 1.35
    width = 4.5
    height = 5.05

    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_rect(slide, left, top, width, 0.5, TEAL, TEAL)
    add_text(slide, left + 0.2, top + 0.13, width - 0.4, 0.25, "Minimum Product Spec", size=12, bold=True, color=WHITE)

    add_rect(slide, left + 0.2, top + 0.58, width - 0.4, 0.45, RGBColor(232, 244, 246), RGBColor(175, 205, 210))
    add_text(slide, left + 0.32, top + 0.66, width - 0.64, 0.14, "NSW Pattern Book Precedent", size=9, bold=True, color=TEAL)
    add_text(slide, left + 0.32, top + 0.82, width - 0.64, 0.16, "Reusable template + controlled site adaptation", size=8, color=SUB)

    specs = [
        "Dwelling type: Detached home",
        "Template basis: NSW Pattern Book-style reusable design",
        "Internal area target: Approx 140-160 sqm",
        "Bedroom count: 3",
        "Bathroom count: 2",
        "Living zones: 2",
        "Parking: Single garage + driveway bay",
        "Outdoor space: Private backyard",
        "Accessibility baseline: Step-free entry where feasible",
        "Compliance: NCC and relevant state code compliant",
    ]

    y = top + 1.08
    for line in specs:
        add_text(slide, left + 0.2, y, width - 0.4, 0.3, "• " + line, size=9, color=TITLE)
        y += 0.3

    # Numeric chips
    chip_w = 1.28
    chip_y = top + height - 0.8
    for i, txt in enumerate(["3-2-1 + drive", "Backyard lot", "Repeatable"]):
        chip_x = left + 0.2 + i * (chip_w + 0.12)
        add_rect(slide, chip_x, chip_y, chip_w, 0.42, RGBColor(255, 255, 255), RGBColor(205, 212, 220))
        add_text(slide, chip_x, chip_y + 0.12, chip_w, 0.2, txt, size=8, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    add_text(
        slide,
        0.45,
        0.16,
        12.2,
        0.52,
        "Product Definition: What Is Being Delivered",
        size=30,
        bold=True,
        color=TITLE,
    )
    add_text(
        slide,
        0.45,
        0.74,
        12.2,
        0.45,
        "NSW Pattern Book precedent: fixed minimum specs with controlled site adaptation, private backyard, and integrated parking.",
        size=12,
        color=SUB,
    )

    draw_floorplan_panel(slide)
    draw_spec_panel(slide)

    add_rect(slide, 0.45, 6.52, 12.4, 0.6, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.72,
        6.71,
        12.0,
        0.25,
        "NSW Pattern Book logic applied: defined product standards turn policy intent into a buildable, auditable housing unit.",
        size=12,
        bold=True,
        color=WHITE,
    )

    add_text(
        slide,
        0.45,
        7.16,
        12.2,
        0.2,
        "Pattern Book learnings from research/nsw.txt plus income_title_outline.txt and procurement.txt.",
        size=8,
        color=SUB,
    )

    notes = (
        "This slide answers a simple question: what exactly are we delivering?\n\n"
        "Income Title is not a broad subsidy for whatever the market supplies. It is a defined home product with fixed minimum standards.\n\n"
        "The NSW Housing Pattern Book is the practical precedent: reusable templates with controlled adaptation can keep quality consistent while reducing redesign and approvals friction.\n\n"
        "That matters for execution. Standardisation improves procurement comparability, reduces design churn, and helps keep quality and cost inside clear bounds.\n\n"
        "For households, these are normal family homes: three bedrooms, two bathrooms, two living zones, single-storey form, and practical liveability.\n\n"
        "For government and delivery partners, the product is measurable and auditable. That supports faster approvals, cleaner tendering, and reliable scale-up."
    )

    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()
