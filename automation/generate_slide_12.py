from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-12.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
ORANGE = RGBColor(198, 94, 31)
SLATE = RGBColor(71, 94, 146)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)


def add_text(slide, left, top, width, height, text, size=14, bold=False, color=TITLE, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_rule_card(slide, left, top, width, height, title, detail, accent):
    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_rect(slide, left, top, width, 0.34, accent, accent)
    add_text(slide, left + 0.12, top + 0.09, width - 0.24, 0.18, title, size=10, bold=True, color=WHITE)
    add_text(slide, left + 0.12, top + 0.45, width - 0.24, height - 0.58, detail, size=9, color=TITLE)


def add_arrow(slide, left, top):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(0.55), Inches(0.2))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(160, 167, 176)
    arrow.line.color.rgb = RGBColor(160, 167, 176)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_text(slide, 0.45, 0.16, 12.2, 0.5, "Financial Mechanics for Households", size=30, bold=True)
    add_text(slide, 0.45, 0.74, 12.2, 0.42, "Repayments align to income and avoid mortgage-style compounding debt burden.", size=12, color=SUB)

    add_rule_card(
        slide,
        0.55,
        1.35,
        3.0,
        2.0,
        "Rule 1: Payment Rate",
        "Indicative contribution is 10-12% of household income, aligning payment burden to capacity.",
        TEAL,
    )
    add_rule_card(
        slide,
        3.78,
        1.35,
        3.0,
        2.0,
        "Rule 2: Annual Floor",
        "A minimum annual floor preserves progress while hardship settings can adjust timing.",
        ORANGE,
    )
    add_rule_card(
        slide,
        7.01,
        1.35,
        3.0,
        2.0,
        "Rule 3: Indexed Cap",
        "Completion target is transparent and indexed by predefined policy method.",
        SLATE,
    )
    add_rule_card(
        slide,
        10.24,
        1.35,
        2.6,
        2.0,
        "Rule 4: Conversion",
        "When cap conditions are met, title converts to freehold.",
        CHARCOAL,
    )

    add_rect(slide, 0.55, 3.62, 12.24, 1.7, RGBColor(252, 253, 254), RGBColor(210, 216, 223))
    add_text(slide, 0.8, 3.78, 2.0, 0.2, "Flow", size=10, bold=True, color=SUB)

    steps = [
        ("Income", "Household earnings"),
        ("Contribution", "10-12% payment"),
        ("Progress", "Cumulative completion"),
        ("Conversion", "Freehold transfer"),
    ]
    x = 1.3
    for i, (title, sub) in enumerate(steps):
        add_rect(slide, x, 4.08, 2.3, 0.88, RGBColor(241, 244, 247), RGBColor(205, 212, 220))
        add_text(slide, x + 0.12, 4.24, 2.06, 0.2, title, size=10, bold=True, color=TITLE, align=PP_ALIGN.CENTER)
        add_text(slide, x + 0.12, 4.48, 2.06, 0.18, sub, size=8, color=SUB, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_arrow(slide, x + 2.38, 4.42)
        x += 2.95

    add_rect(slide, 0.45, 6.35, 12.4, 0.75, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.72,
        6.58,
        12.0,
        0.28,
        "Households pay by capacity and progress by rule, not by compounding debt exposure.",
        size=12,
        bold=True,
        color=WHITE,
    )
    add_text(slide, 0.45, 7.16, 12.2, 0.2, "Slide basis: docs/slides/slide-map.md Section 12.", size=8, color=SUB)

    notes = (
        "This slide defines the household contract in four clear rules.\n\n"
        "Payments are linked to income, constrained by floor/cap settings, and completed through an auditable trigger to title conversion.\n\n"
        "The objective is stable progress to ownership with lower stress sensitivity than fixed mortgage servicing."
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()