from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-03-1-march-replay.deck.pptx"
IMAGES_DIR = ROOT / "outputs/images"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
PANEL = RGBColor(248, 250, 252)
BORDER = RGBColor(215, 221, 228)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(198, 94, 31)


def find_background_image() -> Path:
    candidates = sorted(IMAGES_DIR.glob("*11.33.45*watermark.png"))
    if not candidates:
        raise FileNotFoundError("Could not find the requested watermark background image in outputs/images.")
    return candidates[0]


def add_full_slide_background_image(slide, image_path: Path, slide_width, slide_height):
    picture = slide.shapes.add_picture(str(image_path), 0, 0, width=slide_width, height=slide_height)
    # Send the picture behind all other shapes.
    sp_tree = slide.shapes._spTree
    sp_tree.remove(picture._element)
    sp_tree.insert(2, picture._element)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_transcript_card(slide, left, top, width, height, speaker, lines):
    add_rect(slide, left, top, width, height, PANEL)
    add_rect(slide, left, top, 0.12, height, ACCENT, ACCENT)
    add_text(slide, left + 0.24, top + 0.14, width - 0.36, 0.25, speaker, size=10, bold=True, color=SUB)

    y = top + 0.48
    for line in lines:
        add_text(slide, left + 0.26, y, width - 0.44, 0.34, line, size=11, color=TITLE)
        y += 0.34


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_full_slide_background_image(slide, find_background_image(), prs.slide_width, prs.slide_height)

    add_text(
        slide,
        0.45,
        0.16,
        12.2,
        0.52,
        "Replay: March for Australia Voices",
        size=30,
        bold=True,
        color=TITLE,
    )
    add_text(
        slide,
        0.45,
        0.78,
        12.2,
        0.38,
        "Verbatim-style excerpts from participant speech in research notes.",
        size=12,
        color=SUB,
    )

    speaker_a_lines = [
        '"They\'re creating voting blocks in 15 minute cities',
        'where your kids will never see a backyard."',
    ]
    speaker_b_lines = [
        '"I can\'t sell that dream to my workmates, and it\'s not good enough."',
        '"How are you going to pay for an $800,000 mortgage?"',
        '"People are protesting because they feel locked out."',
        '"We can\'t build enough homes to house people in this country."',
    ]

    add_transcript_card(slide, 0.55, 1.45, 5.85, 3.05, "Speaker A", speaker_a_lines)
    add_transcript_card(slide, 6.93, 1.45, 5.85, 3.05, "Speaker B", speaker_b_lines)

    add_rect(slide, 0.55, 4.78, 12.24, 1.42, RGBColor(252, 253, 254), RGBColor(215, 221, 228))
    add_text(slide, 0.84, 4.94, 11.7, 0.25, "Why this replay matters", size=11, bold=True, color=SUB)
    add_text(
        slide,
        0.88,
        5.24,
        11.5,
        0.7,
        "These comments capture a lived-experience narrative: blocked ownership, debt anxiety, family-formation pressure, and distrust in delivery systems.",
        size=12,
        color=TITLE,
    )

    add_rect(slide, 0.45, 6.45, 12.4, 0.7, RGBColor(245, 247, 250), RGBColor(215, 221, 228))
    add_text(
        slide,
        0.72,
        6.67,
        11.9,
        0.26,
        "Replay framing only. Full transcript references are in research/quotes.txt under Save Australia.",
        size=10,
        bold=True,
        color=TITLE,
    )

    notes = (
        "This slide is a standalone replay moment between the quote wall and technical diagnosis.\n\n"
        "Use it to let audience members hear the emotional and practical stress in direct language before transitioning back to data and mechanism.\n\n"
        "If needed, acknowledge that speech fragments can be compressed for time, while the full transcript remains in the research file."
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()