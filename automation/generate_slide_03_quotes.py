from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-03-quotes.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
ORANGE = RGBColor(198, 94, 31)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_quote_card(slide, left, top, width, height, quote, speaker, source, accent):
    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_rect(slide, left, top, 0.12, height, accent, accent)
    add_text(slide, left + 0.28, top + 0.18, width - 0.5, 0.3, '"', size=28, bold=True, color=accent)
    add_text(slide, left + 0.42, top + 0.22, width - 0.72, 1.1, quote, size=20, bold=True, color=TITLE)
    add_text(slide, left + 0.42, top + 1.48, width - 0.72, 0.22, speaker, size=9, bold=True, color=SUB)
    add_text(slide, left + 0.42, top + 1.72, width - 0.72, 0.2, source, size=8, color=SUB)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    add_text(
        slide,
        0.45,
        0.16,
        12.2,
        0.52,
        "Human Consequences: What People Say the Crisis Feels Like",
        size=28,
        bold=True,
        color=TITLE,
    )
    add_text(
        slide,
        0.45,
        0.74,
        12.2,
        0.42,
        "Housing stress is not abstract. It changes whether people can form households, plan families, or believe ownership is still possible.",
        size=12,
        color=SUB,
    )

    add_quote_card(
        slide,
        0.55,
        1.45,
        4.0,
        2.2,
        "We’ve got no hope really of being able to afford a place of our own.",
        "Young Australian couple interviewed on housing affordability",
        "Source: ABC News",
        TEAL,
    )
    add_quote_card(
        slide,
        4.67,
        1.45,
        4.0,
        2.2,
        "There is a huge issue… where housing is seen as an asset, not a human right.",
        "Interviewee describing systemic housing problems",
        "Source: ABC News",
        ORANGE,
    )
    add_quote_card(
        slide,
        8.79,
        1.45,
        4.0,
        2.2,
        "Renters and owner-occupiers are being squeezed by cost of living and the housing crisis.",
        "Student and union housing campaign representatives",
        "Source: CityHub",
        CHARCOAL,
    )

    add_rect(slide, 0.55, 4.02, 12.24, 1.58, WHITE, RGBColor(210, 216, 223))
    add_text(slide, 0.82, 4.18, 2.4, 0.25, "What these quotes show", size=11, bold=True, color=SUB)
    bullets = [
        "The problem is felt as blocked entry, not just higher prices.",
        "People describe housing as a social contract failure, not a normal market cycle.",
        "The crisis is influencing family formation, life planning, and trust in institutions.",
    ]
    y = 4.48
    for bullet in bullets:
        add_text(slide, 0.88, y, 11.5, 0.22, f"• {bullet}", size=10, color=TITLE)
        y += 0.28

    add_rect(slide, 0.45, 6.55, 12.4, 0.6, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.72,
        6.74,
        12.0,
        0.25,
        "If mainstream workers and renters believe normal housing milestones are gone, the issue is no longer cyclical affordability pressure. It is structural exclusion.",
        size=12,
        bold=True,
        color=WHITE,
    )
    add_text(slide, 0.45, 7.16, 12.2, 0.2, "Source file: research/quotes.txt", size=8, color=SUB)

    notes = (
        "This slide is here to humanise the housing problem before moving into system diagnosis.\n\n"
        "The three quotes were selected because they are short, attributable, and together cover blocked ownership, social meaning, and broad cost-of-living squeeze.\n\n"
        "The point is not that each quote proves the case on its own. The point is that the pattern is recognisable across different speakers and settings: people do not describe a temporary inconvenience, they describe a failure of normal life progression.\n\n"
        "That is why the next section shifts from sentiment to mechanics: what in the housing system is producing this outcome?"
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()