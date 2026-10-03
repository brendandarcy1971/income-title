from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-11.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
ORANGE = RGBColor(198, 94, 31)
SLATE = RGBColor(71, 94, 146)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_bullets(slide, left, top, width, height, lines, size=8, color=TITLE):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_stage_card(slide, left, top, width, title, details, output_text, accent):
    add_rect(slide, left, top, width, 0.56, accent, accent)
    add_text(slide, left + 0.12, top + 0.17, width - 0.24, 0.2, title, size=9, bold=True, color=WHITE)
    add_rect(slide, left, top + 0.56, width, 1.0, RGBColor(255, 255, 255), RGBColor(210, 216, 223))
    add_text(slide, left + 0.1, top + 0.62, width - 0.2, 0.34, details, size=8, color=TITLE)
    add_text(slide, left + 0.1, top + 0.95, width - 0.2, 0.5, f"Output: {output_text}", size=7, color=SUB)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    add_text(slide, 0.45, 0.14, 12.2, 0.44, "Operating Cycle: 100,000 Lots, End-to-End", size=24, bold=True)
    add_text(
        slide,
        0.45,
        0.58,
        12.2,
        0.36,
        "State-led land control + repeatable tranches convert policy intent into measurable housing throughput.",
        size=11,
        color=SUB,
    )

    # Stage timeline with outputs
    add_rect(slide, 0.45, 1.0, 12.42, 2.28, PANEL, BORDER)
    add_text(slide, 0.62, 1.13, 5.0, 0.22, "7-Stage Tranche Engine (with output gates)", size=10, bold=True, color=SUB)

    stages = [
        ("1. Declare", "Precinct declaration", "Gazetted corridor", TEAL),
        ("2. Acquire", "Acquisition + settlement", "Compensation finalized", ORANGE),
        ("3. Approve", "Code approvals", "Subdivision consent", SLATE),
        ("4. Service", "Civil + utility works", "Serviced-lot certificate", TEAL),
        ("5. Release", "Lot allocation window", "Awarded lot register", ORANGE),
        ("6. Build", "Vertical build + QA", "Practical completion", SLATE),
        ("7. Close", "Handover + audit", "Tranche close report", CHARCOAL),
    ]
    x = 0.62
    for title, detail, output_text, color in stages:
        add_stage_card(slide, x, 1.34, 1.67, title, detail, output_text, color)
        x += 1.74

    # Lower left: responsibility split
    add_rect(slide, 0.45, 3.42, 7.95, 2.72, PANEL, BORDER)
    add_text(slide, 0.62, 3.55, 3.8, 0.2, "Responsibility Split", size=10, bold=True, color=SUB)

    add_rect(slide, 0.62, 3.84, 1.8, 2.12, RGBColor(255, 255, 255), RGBColor(215, 220, 226))
    add_text(slide, 0.72, 3.95, 1.55, 0.2, "Commonwealth", size=8, bold=True, color=SLATE)
    add_bullets(slide, 0.72, 4.16, 1.55, 1.72, ["Funding envelope", "National standards", "Delivery-linked incentives"], size=7)

    add_rect(slide, 2.52, 3.84, 1.9, 2.12, RGBColor(255, 255, 255), RGBColor(215, 220, 226))
    add_text(slide, 2.64, 3.95, 1.65, 0.2, "State Government", size=8, bold=True, color=ORANGE)
    add_bullets(slide, 2.64, 4.16, 1.65, 1.72, ["Acquisition authority", "Planning + subdivision", "Utility/civil coordination", "Package contracting"], size=7)

    add_rect(slide, 4.52, 3.84, 1.8, 2.12, RGBColor(255, 255, 255), RGBColor(215, 220, 226))
    add_text(slide, 4.64, 3.95, 1.55, 0.2, "Contractors", size=8, bold=True, color=TEAL)
    add_bullets(slide, 4.64, 4.16, 1.55, 1.72, ["Survey/engineering", "Civil execution", "Vertical build", "Defects closeout"], size=7)

    add_rect(slide, 6.42, 3.84, 1.8, 2.12, RGBColor(255, 255, 255), RGBColor(215, 220, 226))
    add_text(slide, 6.52, 3.95, 1.55, 0.2, "Assurance", size=8, bold=True, color=CHARCOAL)
    add_bullets(slide, 6.52, 4.16, 1.55, 1.72, ["Cost audit", "Quality audit", "Schedule verification", "Exception triggers"], size=7)

    # Lower right: program math, cadence, and targeted eligibility
    add_rect(slide, 8.55, 3.42, 4.32, 2.72, PANEL, BORDER)
    add_text(slide, 8.74, 3.55, 2.8, 0.2, "Program Controls + Throughput", size=10, bold=True, color=SUB)

    add_rect(slide, 8.74, 3.84, 2.05, 1.06, RGBColor(255, 255, 255), RGBColor(215, 220, 226))
    add_text(slide, 8.86, 3.94, 1.8, 0.18, "Tranche Math", size=8, bold=True, color=SLATE)
    add_text(slide, 8.86, 4.13, 1.8, 0.7, "100,000 lots total\n10 tranches x 10,000\nLaunch every ~6 months", size=7, color=TITLE)

    add_rect(slide, 10.92, 3.84, 1.75, 1.06, RGBColor(255, 255, 255), RGBColor(215, 220, 226))
    add_text(slide, 11.02, 3.94, 1.55, 0.18, "Steady State", size=8, bold=True, color=TEAL)
    add_text(slide, 11.02, 4.13, 1.55, 0.7, "3-4 active tranches\nParallel civil/build\nMonthly KPI review", size=7, color=TITLE)

    add_rect(slide, 8.74, 5.02, 3.93, 1.12, RGBColor(232, 244, 246), RGBColor(175, 205, 210))
    add_text(slide, 8.88, 5.12, 2.4, 0.16, "Targeted Eligibility Overlay", size=8, bold=True, color=TEAL)
    add_bullets(
        slide,
        8.88,
        5.28,
        3.66,
        0.78,
        [
            "Base test: no current ownership, long-term renting, income-band cap",
            "Priority score: rent stress, years renting, household need",
            "Safeguards: anti-gaming checks + independent review",
        ],
        size=7,
    )

    # KPI strip and close
    add_rect(slide, 0.45, 6.24, 12.42, 0.64, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.6,
        6.34,
        12.1,
        0.16,
        "KPI dashboard: cost/serviced lot, cost/completed dwelling, approval cycle time, lots delivered/month, on-time handover, 90-day defects.",
        size=8,
        color=WHITE,
    )
    add_text(
        slide,
        0.6,
        6.54,
        12.1,
        0.16,
        "Execution principle: state land control captures uplift, private delivery panels preserve speed and contestability.",
        size=8,
        color=WHITE,
    )
    add_text(slide, 0.45, 6.94, 12.2, 0.2, "Slide basis: docs/slides/slide-11.plan.md", size=8, color=SUB)

    notes = (
        "This slide defines the full operating cadence for delivering 100,000 lots via overlapping 10,000-lot tranches.\n\n"
        "Each stage has an explicit output gate, and each actor has auditable responsibilities with monthly KPI review.\n\n"
        "Targeted eligibility is presented as an operational control to prioritize households excluded from ownership while protecting fairness and program legitimacy."
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()