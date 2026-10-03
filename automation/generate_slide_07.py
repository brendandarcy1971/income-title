from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-07.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(24, 24, 24)
SUB = RGBColor(78, 78, 78)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(27, 92, 100)
ORANGE = RGBColor(198, 94, 31)
SLATE = RGBColor(70, 95, 145)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(218, 224, 229)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=12,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_bullets(slide, left, top, width, height, items, size=8, color=TITLE):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_stage(slide, left, title, bullets, accent):
    add_rect(slide, left, 1.72, 3.95, 3.7, WHITE, BORDER)
    add_rect(slide, left, 1.72, 3.95, 0.58, accent, accent)
    add_text(slide, left + 0.14, 1.89, 3.7, 0.2, title, size=11, bold=True, color=WHITE)
    add_bullets(slide, left + 0.18, 2.45, 3.6, 2.7, bullets, size=9)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_text(slide, 0.45, 0.16, 12.3, 0.44, "Income Title Lifecycle: Enter, Repay by Income, Convert to Freehold", size=24, bold=True)
    add_text(
        slide,
        0.45,
        0.62,
        12.3,
        0.34,
        "Occupancy starts immediately. Ownership is completed through structured income-linked progression.",
        size=11,
        color=SUB,
    )

    add_stage(
        slide,
        0.65,
        "1. Entry",
        [
            "Eligible household allocated compliant home",
            "No deposit barrier at entry",
            "Immediate occupancy rights begin",
            "Clear acceptance and onboarding rules",
        ],
        TEAL,
    )
    add_stage(
        slide,
        4.7,
        "2. Income-Linked Progression",
        [
            "Annual contribution set as share of income",
            "Compliance and reporting rules apply",
            "Mobility options for transfer/sale before completion",
            "Payment burden tracks earning capacity",
        ],
        SLATE,
    )
    add_stage(
        slide,
        8.75,
        "3. Title Conversion",
        [
            "Threshold reached under scheme rules",
            "Title converts to full freehold",
            "Scheme balance closes for dwelling",
            "Household exits with standard ownership",
        ],
        ORANGE,
    )

    # Connectors
    add_text(slide, 4.28, 3.36, 0.35, 0.25, "→", size=28, bold=True, color=SUB, align=PP_ALIGN.CENTER)
    add_text(slide, 8.33, 3.36, 0.35, 0.25, "→", size=28, bold=True, color=SUB, align=PP_ALIGN.CENTER)

    # Rule bar and fairness panel
    add_rect(slide, 0.65, 5.58, 8.05, 1.08, RGBColor(232, 244, 246), RGBColor(175, 205, 210))
    add_text(slide, 0.82, 5.72, 7.7, 0.2, "Rule Bar", size=9, bold=True, color=TEAL)
    add_text(
        slide,
        0.82,
        5.95,
        7.7,
        0.48,
        "Clear entry rules. Predictable repayment rules. Automatic conversion rules.",
        size=10,
        bold=True,
        color=TITLE,
    )

    add_rect(slide, 8.9, 5.58, 3.8, 1.08, PANEL, BORDER)
    add_text(slide, 9.08, 5.72, 3.4, 0.2, "Policy Framing", size=9, bold=True, color=SUB)
    add_text(
        slide,
        9.08,
        5.95,
        3.4,
        0.46,
        "Targeted to households locked out of ownership, not a universal demand subsidy.",
        size=8,
        color=TITLE,
    )

    add_text(slide, 0.45, 6.94, 12.1, 0.2, "Slide basis: docs/slides/slides-06-07.detailed.md", size=8, color=SUB)

    notes = (
        "This is the complete lifecycle in one visual: entry, progression, conversion.\n\n"
        "The policy intent is clarity for households and predictability for program administration.\n\n"
        "Link this slide directly after policy precedent material to show how the model works in practice."
    )
    nf = slide.notes_slide.notes_text_frame
    nf.clear()
    nf.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()
