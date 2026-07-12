from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-09.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
ORANGE = RGBColor(198, 94, 31)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def draw_cost_chart(slide, left, top, width, height):
    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_text(slide, left + 0.16, top + 0.08, width - 0.32, 0.2, "Target Cost Stack Per Dwelling", size=10, bold=True, color=SUB)

    chart_left = left + 0.45
    chart_top = top + 0.45
    chart_w = width - 0.9
    chart_h = height - 1.15

    total = 350
    land = 100
    build = 250

    bar_w = 2.0
    x = chart_left + (chart_w - bar_w) / 2
    base_y = chart_top + chart_h

    land_h = chart_h * (land / total)
    build_h = chart_h * (build / total)

    # Build layer
    add_rect(slide, x, base_y - build_h, bar_w, build_h, TEAL, TEAL)
    add_text(slide, x + 0.2, base_y - (build_h / 2), 1.6, 0.22, "$250k\nBuild", size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Land layer
    add_rect(slide, x, base_y - build_h - land_h, bar_w, land_h, ORANGE, ORANGE)
    add_text(slide, x + 0.2, base_y - build_h - (land_h / 2), 1.6, 0.22, "$100k\nServiced land", size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Total label
    add_text(slide, x + 0.1, base_y - build_h - land_h - 0.22, 1.8, 0.18, "Total target: $350k", size=9, bold=True, color=SUB, align=PP_ALIGN.CENTER)

    # Side assumptions
    add_rect(slide, left + width - 2.35, top + 0.45, 2.05, height - 1.0, RGBColor(255, 255, 255), RGBColor(210, 216, 223))
    add_text(slide, left + width - 2.18, top + 0.55, 1.7, 0.2, "Assumptions", size=9, bold=True, color=SUB)
    items = [
        "Standardised product spec",
        "Repeatable procurement lots",
        "Serviced land discipline",
        "Multi-contractor competition",
        "QA and compliance controls",
    ]
    y = top + 0.78
    for item in items:
        add_text(slide, left + width - 2.16, y, 1.72, 0.25, "• " + item, size=8, color=TITLE)
        y += 0.3


def draw_threshold_panel(slide, left, top, width, height):
    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_text(slide, left + 0.16, top + 0.08, width - 0.32, 0.2, "Target Economics Checks", size=10, bold=True, color=SUB)

    checks = [
        ("Design target", "$350k total delivered cost"),
        ("Land envelope", "~$100k serviced land"),
        ("Build envelope", "~$250k construction"),
        ("Guardrail", "Template + procurement discipline"),
    ]
    y = top + 0.42
    for label, value in checks:
        add_rect(slide, left + 0.16, y, width - 0.32, 0.4, RGBColor(255, 255, 255), RGBColor(210, 216, 223))
        add_text(slide, left + 0.28, y + 0.08, 1.45, 0.15, label, size=8, bold=True, color=SUB)
        add_text(slide, left + 1.78, y + 0.08, width - 2.1, 0.15, value, size=9, color=TITLE)
        y += 0.48


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    add_text(slide, 0.45, 0.16, 12.2, 0.5, "Cost Structure and Target Economics", size=30, bold=True)
    add_text(
        slide,
        0.45,
        0.74,
        12.2,
        0.44,
        "The $350k target is a designed outcome, not a hope: indicative split is ~$100k serviced land + ~$250k build.",
        size=12,
        color=SUB,
    )

    draw_cost_chart(slide, left=0.55, top=1.35, width=7.85, height=4.95)
    draw_threshold_panel(slide, left=8.55, top=1.35, width=4.3, height=4.95)

    add_rect(slide, 0.45, 6.42, 12.4, 0.7, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.72,
        6.65,
        12.0,
        0.25,
        "Target economics are achieved by product standardisation and procurement design, not by assuming market prices will self-correct.",
        size=11,
        bold=True,
        color=WHITE,
    )
    add_text(slide, 0.45, 7.16, 12.2, 0.2, "Slide basis: docs/slides/slide-map.md (Section 9 target split).", size=8, color=SUB)

    notes = (
        "This slide anchors the cost claim in a simple designed stack.\n\n"
        "The key point is that 350k is not treated as optimism. It is treated as a governed target composed of two envelopes: serviced land and build delivery.\n\n"
        "The assumptions panel states the operating controls needed to keep cost inside that envelope: standard product, repeatable lots, and competitive delivery under compliance rules.\n\n"
        "Transition to procurement architecture next: the mechanism that turns this cost design into executable contracts."
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()