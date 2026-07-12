from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT / "outputs/decks/slides-06-07.deck.pptx"


# Color palette
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(20, 20, 20)
CHARCOAL = RGBColor(54, 58, 66)
LIGHT_BG = RGBColor(246, 248, 250)
HECS_TEAL = RGBColor(0, 102, 102)
PBS_ORANGE = RGBColor(198, 94, 31)
BRIDGE_GREY = RGBColor(75, 78, 85)
CHIP_BG = RGBColor(235, 239, 243)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False, color=BLACK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, lines, font_size=14, color=BLACK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
    return box


def add_card(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    return shape


def add_chip(slide, left, top, width, height, line1, line2):
    chip = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    chip.fill.solid()
    chip.fill.fore_color.rgb = CHIP_BG
    chip.line.color.rgb = RGBColor(210, 216, 223)

    tf = chip.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = line1
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.font.color.rgb = BLACK

    p2 = tf.add_paragraph()
    p2.text = line2
    p2.font.size = Pt(10)
    p2.font.color.rgb = RGBColor(70, 70, 70)
    return chip


def add_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    notes_frame = notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes_text


def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_BG)

    add_textbox(slide, 0.4, 0.2, 12.5, 0.6, "Proven Policy Logic, New Housing Application", 30, True)
    add_textbox(
        slide,
        0.4,
        0.8,
        12.5,
        0.5,
        "Income Title combines HECS repayment design with PBS procurement discipline.",
        14,
        False,
        RGBColor(60, 60, 60),
    )

    # Left card (HECS)
    add_card(slide, 0.5, 1.5, 4.0, 4.6, HECS_TEAL)
    add_textbox(slide, 0.75, 1.8, 3.5, 0.4, "HECS Principles", 18, True, WHITE)
    add_textbox(slide, 0.75, 2.15, 3.5, 0.3, "Repayment Design", 11, False, WHITE)

    add_bullets(
        slide,
        0.8,
        2.5,
        3.4,
        2.4,
        [
            "Access first, repayment later.",
            "Repayment tied to capacity to pay.",
            "Automatic collection architecture lowers friction.",
            "Public objective with private life outcomes.",
        ],
        12,
        WHITE,
    )

    add_chip(slide, 0.75, 5.0, 3.5, 0.5, "Since 1989", "Income-contingent model in operation")
    add_chip(slide, 0.75, 5.55, 3.5, 0.5, "26% -> 34%", "Bachelor degree+ attainment, 2016 to 2025")

    # Center bridge
    add_card(slide, 4.9, 1.8, 2.4, 4.0, BRIDGE_GREY)
    add_textbox(slide, 5.1, 2.0, 2.0, 0.4, "Transfer to", 14, True, WHITE)
    add_textbox(slide, 5.1, 2.35, 2.0, 0.4, "Housing", 14, True, WHITE)

    add_bullets(
        slide,
        5.1,
        2.9,
        2.0,
        2.6,
        [
            "Access now",
            "Pay by income",
            "Standard product",
            "Structured purchasing",
        ],
        11,
        WHITE,
    )

    # Right card (PBS)
    add_card(slide, 7.7, 1.5, 4.0, 4.6, PBS_ORANGE)
    add_textbox(slide, 7.95, 1.8, 3.5, 0.4, "PBS Principles", 18, True, WHITE)
    add_textbox(slide, 7.95, 2.15, 3.5, 0.3, "Procurement Design", 11, False, WHITE)

    add_bullets(
        slide,
        8.0,
        2.5,
        3.4,
        2.4,
        [
            "Government sets purchasing framework.",
            "Standard specifications reduce fragmentation.",
            "Price discipline via rules and disclosure.",
            "Scale buying power improves value for money.",
        ],
        12,
        WHITE,
    )

    add_chip(slide, 7.95, 5.0, 3.5, 0.5, "226.0 million", "Annual subsidised prescriptions, 2024-25")
    add_chip(slide, 7.95, 5.55, 3.5, 0.5, "$19.1 billion", "Government PBS expenditure, 2024-25")

    add_textbox(
        slide,
        0.5,
        6.6,
        11.2,
        0.35,
        "Sources: PBS Expenditure and Prescriptions Reports 2023-24 and 2024-25; ABS Education and Work 2025.",
        9,
        False,
        RGBColor(95, 95, 95),
    )

    notes = (
        "Australia already runs complex social systems with hard fiscal discipline.\n\n"
        "HECS demonstrates that income-linked repayment can preserve access while managing burden over time. "
        "PBS demonstrates that a clear purchasing framework with transparent rules can sustain long-term price discipline.\n\n"
        "Income Title combines those logics. From HECS, household repayment scales with capacity. "
        "From PBS, system delivery uses rules, benchmarking, and disciplined procurement.\n\n"
        "This is why the model is credible: it adapts proven Australian mechanisms rather than inventing an untested structure from scratch."
    )
    add_notes(slide, notes)


def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_BG)

    add_textbox(slide, 0.4, 0.2, 12.5, 0.6, "Income Title Lifecycle: Enter, Repay by Income, Convert to Freehold", 28, True)
    add_textbox(
        slide,
        0.4,
        0.8,
        12.5,
        0.5,
        "Occupancy starts immediately. Ownership is completed through structured income-linked repayment.",
        14,
        False,
        RGBColor(60, 60, 60),
    )

    # Lifecycle nodes
    node_w = 3.4
    node_h = 2.4
    y = 1.8

    # Stage 1
    add_card(slide, 0.6, y, node_w, node_h, HECS_TEAL)
    add_textbox(slide, 0.8, y + 0.2, 3.0, 0.35, "1. Entry", 16, True, WHITE)
    add_bullets(
        slide,
        0.85,
        y + 0.6,
        3.0,
        1.6,
        [
            "Eligible household allocated compliant home.",
            "No deposit barrier.",
            "Immediate occupancy rights begin.",
        ],
        11,
        WHITE,
    )

    # Stage 2
    add_card(slide, 4.7, y, node_w, node_h, BRIDGE_GREY)
    add_textbox(slide, 4.9, y + 0.2, 3.0, 0.35, "2. Income-Linked Progression", 15, True, WHITE)
    add_bullets(
        slide,
        4.95,
        y + 0.6,
        3.0,
        1.6,
        [
            "Annual contribution set as share of household income.",
            "Standard administration and compliance rules.",
            "Mobility pathways for transfer or sale.",
        ],
        11,
        WHITE,
    )

    # Stage 3
    add_card(slide, 8.8, y, node_w, node_h, PBS_ORANGE)
    add_textbox(slide, 9.0, y + 0.2, 3.0, 0.35, "3. Title Conversion", 16, True, WHITE)
    add_bullets(
        slide,
        9.05,
        y + 0.6,
        3.0,
        1.6,
        [
            "When repayment threshold is reached, title converts.",
            "Household holds full freehold ownership.",
            "Scheme balance closes for that dwelling.",
        ],
        11,
        WHITE,
    )

    # Connectors
    arrow1 = slide.shapes.add_shape(33, Inches(4.05), Inches(2.7), Inches(0.5), Inches(0.25))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = CHARCOAL
    arrow1.line.color.rgb = CHARCOAL

    arrow2 = slide.shapes.add_shape(33, Inches(8.15), Inches(2.7), Inches(0.5), Inches(0.25))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = CHARCOAL
    arrow2.line.color.rgb = CHARCOAL

    # Numeric callouts cluster
    add_chip(slide, 0.8, 4.5, 2.5, 0.5, "$350k", "Target delivered cost per home")
    add_chip(slide, 3.45, 4.5, 2.7, 0.5, "$100k + $250k", "Indicative land/build split")
    add_chip(slide, 6.3, 4.5, 2.2, 0.5, "10-12%", "Payment rate of income")
    add_chip(slide, 8.65, 4.5, 2.2, 0.5, "$8k-$10k", "Minimum annual payment")
    add_chip(slide, 10.95, 4.5, 1.8, 0.5, "$300k", "Indexed cap trigger")

    # Bottom rule bar
    bar = slide.shapes.add_shape(1, Inches(0.6), Inches(5.35), Inches(12.2), Inches(0.7))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CHARCOAL
    bar.line.color.rgb = CHARCOAL
    add_textbox(
        slide,
        1.0,
        5.55,
        11.6,
        0.3,
        "Clear entry rules. Predictable repayment rules. Automatic conversion rules.",
        14,
        True,
        WHITE,
    )

    notes = (
        "At entry, the household gets immediate occupancy without a deposit hurdle.\n\n"
        "During progression, payments are income-linked rather than fixed mortgage servicing, "
        "improving resilience through income variability.\n\n"
        "At conversion, once the repayment threshold is met under scheme rules, title converts to full freehold.\n\n"
        "This creates clarity for households, auditable administration for government, and repeatable delivery logic for scale."
    )
    add_notes(slide, notes)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_slide_6(prs)
    build_slide_7(prs)

    prs.save(str(OUTPUT_PATH))


if __name__ == "__main__":
    main()
