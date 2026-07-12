from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-14.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
ORANGE = RGBColor(198, 94, 31)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)


def add_text(slide, left, top, width, height, text, size=14, bold=False, color=TITLE, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def draw_curve_panel(slide, left, top, width, height):
    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_text(slide, left + 0.16, top + 0.08, width - 0.32, 0.2, "Comparative Completion Curves", size=10, bold=True, color=SUB)

    chart_left = left + 0.45
    chart_top = top + 0.4
    chart_w = width - 0.9
    chart_h = height - 0.95
    add_rect(slide, chart_left, chart_top, chart_w, chart_h, RGBColor(255, 255, 255), RGBColor(210, 216, 223))

    # Axes
    add_rect(slide, chart_left + 0.45, chart_top + 0.2, 0.01, chart_h - 0.5, RGBColor(170, 176, 184), RGBColor(170, 176, 184))
    add_rect(slide, chart_left + 0.45, chart_top + chart_h - 0.3, chart_w - 0.8, 0.01, RGBColor(170, 176, 184), RGBColor(170, 176, 184))

    # Middle-income curve (teal)
    middle = [
        (chart_left + 0.65, chart_top + chart_h - 0.42),
        (chart_left + 1.6, chart_top + chart_h - 0.95),
        (chart_left + 2.7, chart_top + chart_h - 1.58),
        (chart_left + 3.9, chart_top + chart_h - 2.3),
        (chart_left + 5.1, chart_top + chart_h - 2.95),
    ]
    for i in range(len(middle) - 1):
        x1, y1 = middle[i]
        x2, y2 = middle[i + 1]
        add_rect(slide, x1, min(y1, y2), max(0.08, x2 - x1), max(0.05, abs(y2 - y1) + 0.05), TEAL, TEAL)

    # Higher-income curve (orange)
    higher = [
        (chart_left + 0.65, chart_top + chart_h - 0.42),
        (chart_left + 1.4, chart_top + chart_h - 1.15),
        (chart_left + 2.2, chart_top + chart_h - 1.95),
        (chart_left + 3.1, chart_top + chart_h - 2.95),
        (chart_left + 4.1, chart_top + chart_h - 3.6),
        (chart_left + 5.0, chart_top + chart_h - 4.0),
    ]
    for i in range(len(higher) - 1):
        x1, y1 = higher[i]
        x2, y2 = higher[i + 1]
        add_rect(slide, x1, min(y1, y2), max(0.08, x2 - x1), max(0.05, abs(y2 - y1) + 0.05), ORANGE, ORANGE)

    add_text(slide, chart_left + 4.8, chart_top + chart_h - 2.95, 1.4, 0.16, "Middle-income", size=8, bold=True, color=TEAL)
    add_text(slide, chart_left + 4.7, chart_top + chart_h - 3.95, 1.4, 0.16, "Higher-income", size=8, bold=True, color=ORANGE)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_text(slide, 0.45, 0.16, 12.2, 0.5, "Worked Example: Middle and Higher Income Pathways", size=28, bold=True)
    add_text(slide, 0.45, 0.74, 12.2, 0.42, "Same rules, different earning trajectories, different completion timing.", size=12, color=SUB)

    draw_curve_panel(slide, 0.55, 1.35, 8.15, 4.95)

    # Right-side outcome panel
    add_rect(slide, 8.9, 1.35, 3.95, 4.95, PANEL, BORDER)
    add_text(slide, 9.1, 1.5, 3.5, 0.2, "Completion Ranges", size=10, bold=True, color=SUB)

    add_rect(slide, 9.1, 1.82, 3.55, 1.25, RGBColor(255, 255, 255), RGBColor(205, 212, 220))
    add_text(slide, 9.28, 1.95, 3.2, 0.18, "Middle-income pathway", size=9, bold=True, color=TEAL)
    add_text(slide, 9.28, 2.2, 3.2, 0.18, "Indicative completion:", size=8, color=SUB)
    add_text(slide, 9.28, 2.42, 3.2, 0.2, "~14-17 years", size=12, bold=True, color=TITLE)

    add_rect(slide, 9.1, 3.22, 3.55, 1.25, RGBColor(255, 255, 255), RGBColor(205, 212, 220))
    add_text(slide, 9.28, 3.35, 3.2, 0.18, "Higher-income pathway", size=9, bold=True, color=ORANGE)
    add_text(slide, 9.28, 3.6, 3.2, 0.18, "Indicative completion:", size=8, color=SUB)
    add_text(slide, 9.28, 3.82, 3.2, 0.2, "~10-13 years", size=12, bold=True, color=TITLE)

    add_rect(slide, 9.1, 4.62, 3.55, 1.45, RGBColor(255, 255, 255), RGBColor(205, 212, 220))
    add_text(slide, 9.28, 4.77, 3.2, 0.2, "Interpretation", size=9, bold=True, color=SUB)
    add_text(
        slide,
        9.28,
        5.02,
        3.2,
        0.92,
        "The framework is stable across cohorts. Time-to-completion changes with earnings capacity, not with separate policy rulebooks.",
        size=8,
        color=TITLE,
    )

    add_rect(slide, 0.45, 6.35, 12.4, 0.75, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.72,
        6.58,
        12.0,
        0.28,
        "Income-linked mechanics scale naturally: higher earnings accelerate completion while rule consistency is preserved.",
        size=12,
        bold=True,
        color=WHITE,
    )
    add_text(slide, 0.45, 7.16, 12.2, 0.2, "Slide basis: docs/slides/slide-map.md Section 14.", size=8, color=SUB)

    notes = (
        "This comparison demonstrates capacity-based acceleration under one consistent rule set.\n\n"
        "Higher income does not require different mechanics; it changes contribution pace and therefore time-to-completion.\n\n"
        "That preserves policy simplicity while reflecting real household diversity."
    )
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()