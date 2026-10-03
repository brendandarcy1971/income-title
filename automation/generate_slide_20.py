from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs/decks/slide-20.deck.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(22, 22, 22)
SUB = RGBColor(78, 78, 78)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(53, 58, 66)
TEAL = RGBColor(27, 92, 100)
ORANGE = RGBColor(198, 94, 31)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(218, 224, 229)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=12,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_bullets(slide, left, top, width, height, items, size=8, color=TITLE):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.level = 0
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_panel(slide, left, title, accent, what_changes, what_stays):
    add_rect(slide, left, 1.64, 5.95, 3.95, WHITE, BORDER)
    add_rect(slide, left, 1.64, 5.95, 0.62, accent, accent)
    add_text(slide, left + 0.18, 1.84, 5.6, 0.2, title, size=11, bold=True, color=WHITE)

    add_text(slide, left + 0.2, 2.42, 2.7, 0.2, "What Changes", size=9, bold=True, color=SUB)
    add_bullets(slide, left + 0.2, 2.62, 2.72, 1.5, what_changes, size=8)

    add_text(slide, left + 3.0, 2.42, 2.7, 0.2, "What Stays Private", size=9, bold=True, color=SUB)
    add_bullets(slide, left + 3.0, 2.62, 2.72, 1.5, what_stays, size=8)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_text(slide, 0.45, 0.16, 12.3, 0.46, "Will Private Developers Exit? No - They Shift to Scaled Delivery", size=24, bold=True)
    add_text(
        slide,
        0.45,
        0.62,
        12.3,
        0.36,
        "Government stabilises land and approvals throughput; private firms still deliver civil works and homes at scale.",
        size=11,
        color=SUB,
    )

    add_panel(
        slide,
        0.62,
        "Before: Speculative Bottleneck Model",
        ORANGE,
        [
            "Landbank and rezoning windfall dependence",
            "Approvals volatility and holding-cost risk",
            "Irregular pipeline and workforce swings",
            "Margin quality tied to timing luck",
        ],
        [
            "Survey and planning capability",
            "Civil and utility execution",
            "Home construction operations",
            "Subcontractor ecosystem",
        ],
    )

    add_panel(
        slide,
        6.75,
        "After: Programmed Tranche Delivery",
        TEAL,
        [
            "Predictable 10k-lot tranche cadence",
            "Open tenders and package competition",
            "Lower approval and release uncertainty",
            "Production-led rather than windfall-led returns",
        ],
        [
            "Private civil/build delivery",
            "Quality and defects obligations",
            "Commercial innovation in execution",
            "Performance-based rollover opportunities",
        ],
    )

    add_rect(slide, 0.62, 5.74, 8.22, 0.98, PANEL, BORDER)
    add_text(slide, 0.8, 5.86, 2.7, 0.2, "Policy Guardrails", size=9, bold=True, color=SUB)
    add_bullets(
        slide,
        0.8,
        6.06,
        7.9,
        0.55,
        [
            "Open panel model; no single-state-builder monopoly",
            "Tiered package sizing across Tier 1, mid-tier, and SME firms",
            "Minimum private delivery share per tranche + annual market-impact review",
        ],
        size=8,
    )

    add_rect(slide, 8.98, 5.74, 3.72, 0.98, CHARCOAL, CHARCOAL)
    add_text(slide, 9.14, 5.89, 3.42, 0.58, "This is not nationalisation of construction; it is de-risked market delivery at scale.", size=9, bold=True, color=WHITE)

    add_text(slide, 0.45, 6.94, 12.2, 0.2, "Slide basis: docs/slides/slide-20-private-market-reassurance.plan.md", size=8, color=SUB)

    notes = (
        "The private market does not disappear in this model; it shifts from speculative land-margin dependence to contracted delivery certainty.\n\n"
        "Government's role is to stabilise land and approvals throughput. Private firms still execute civil and vertical delivery under competitive procurement.\n\n"
        "Use guardrails to maintain contractor diversity and avoid concentration risk."
    )
    nf = slide.notes_slide.notes_text_frame
    nf.clear()
    nf.text = notes

    prs.save(str(OUT))


if __name__ == "__main__":
    main()
