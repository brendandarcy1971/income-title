from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


OUT = "proposal/income_title_slide_08.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
WHITE = RGBColor(255, 255, 255)
CHARCOAL = RGBColor(54, 58, 66)
TEAL = RGBColor(30, 95, 102)
PANEL = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)
ACCENT = RGBColor(198, 94, 31)


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=TITLE,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_rect(slide, left, top, width, height, fill, line=BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def add_callout(slide, left, top, text):
    chip = add_rect(slide, left, top, 2.0, 0.32, RGBColor(255, 255, 255), BORDER)
    _ = chip
    add_text(slide, left + 0.08, top + 0.05, 1.85, 0.22, text, size=9, bold=True, color=TEAL)


def draw_floorplan_panel(slide):
    left = 0.45
    top = 1.35
    width = 7.75
    height = 5.05

    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_text(slide, left + 0.18, top + 0.08, 5.0, 0.28, "Illustrative Standard Floorplan", size=11, bold=True, color=SUB)

    # Main floorplan frame
    plan_left = left + 0.35
    plan_top = top + 0.48
    plan_w = 6.0
    plan_h = 3.95
    add_rect(slide, plan_left, plan_top, plan_w, plan_h, RGBColor(252, 253, 254), RGBColor(180, 186, 194))

    # Rooms (diagram only)
    add_rect(slide, plan_left + 0.12, plan_top + 0.12, 1.82, 1.22, RGBColor(245, 249, 252), RGBColor(185, 195, 205))
    add_text(slide, plan_left + 0.22, plan_top + 0.58, 1.6, 0.22, "Bedroom 1", size=8, color=SUB)

    add_rect(slide, plan_left + 2.02, plan_top + 0.12, 1.78, 1.22, RGBColor(245, 249, 252), RGBColor(185, 195, 205))
    add_text(slide, plan_left + 2.16, plan_top + 0.58, 1.4, 0.22, "Bedroom 2", size=8, color=SUB)

    add_rect(slide, plan_left + 3.88, plan_top + 0.12, 1.94, 1.22, RGBColor(245, 249, 252), RGBColor(185, 195, 205))
    add_text(slide, plan_left + 4.05, plan_top + 0.58, 1.5, 0.22, "Bedroom 3", size=8, color=SUB)

    add_rect(slide, plan_left + 0.12, plan_top + 1.45, 1.48, 1.0, RGBColor(249, 246, 241), RGBColor(190, 193, 199))
    add_text(slide, plan_left + 0.22, plan_top + 1.84, 1.2, 0.2, "Bath", size=8, color=SUB)

    add_rect(slide, plan_left + 1.68, plan_top + 1.45, 1.48, 1.0, RGBColor(249, 246, 241), RGBColor(190, 193, 199))
    add_text(slide, plan_left + 1.8, plan_top + 1.84, 1.3, 0.2, "Ensuite", size=8, color=SUB)

    add_rect(slide, plan_left + 3.24, plan_top + 1.45, 2.58, 1.0, RGBColor(242, 249, 244), RGBColor(185, 195, 185))
    add_text(slide, plan_left + 3.45, plan_top + 1.84, 2.2, 0.2, "Family Living", size=8, color=SUB)

    add_rect(slide, plan_left + 0.12, plan_top + 2.56, 2.35, 1.34, RGBColor(242, 249, 244), RGBColor(185, 195, 185))
    add_text(slide, plan_left + 0.28, plan_top + 3.1, 2.0, 0.2, "Living Area 2", size=8, color=SUB)

    add_rect(slide, plan_left + 2.55, plan_top + 2.56, 3.27, 1.34, RGBColor(243, 246, 252), RGBColor(185, 195, 205))
    add_text(slide, plan_left + 2.8, plan_top + 3.1, 2.9, 0.2, "Kitchen + Dining", size=8, color=SUB)

    # Callout chips
    add_callout(slide, left + 0.2, top + 4.55, "3 Bedrooms")
    add_callout(slide, left + 2.28, top + 4.55, "2 Bathrooms")
    add_callout(slide, left + 4.36, top + 4.55, "2 Living Areas")

    add_callout(slide, left + 0.2, top + 4.9, "Single-storey layout")
    add_callout(slide, left + 2.28, top + 4.9, "Family kitchen/dining")
    add_callout(slide, left + 4.36, top + 4.9, "Outdoor access + storage")


def draw_spec_panel(slide):
    left = 8.35
    top = 1.35
    width = 4.5
    height = 5.05

    add_rect(slide, left, top, width, height, PANEL, BORDER)
    add_rect(slide, left, top, width, 0.5, TEAL, TEAL)
    add_text(slide, left + 0.2, top + 0.13, width - 0.4, 0.25, "Minimum Product Spec", size=12, bold=True, color=WHITE)

    specs = [
        "Dwelling type: Detached home",
        "Internal area target: Approx 140-160 sqm",
        "Bedroom count: 3",
        "Bathroom count: 2",
        "Living zones: 2",
        "Accessibility baseline: Step-free entry where feasible",
        "Compliance: NCC and relevant state code compliant",
    ]

    y = top + 0.65
    for line in specs:
        add_text(slide, left + 0.2, y, width - 0.4, 0.3, "• " + line, size=9, color=TITLE)
        y += 0.37

    # Numeric chips
    chip_w = 1.28
    chip_y = top + height - 0.8
    for i, txt in enumerate(["3-2-2", "140-160 sqm", "Repeatable"]):
        chip_x = left + 0.2 + i * (chip_w + 0.12)
        add_rect(slide, chip_x, chip_y, chip_w, 0.42, RGBColor(255, 255, 255), RGBColor(205, 212, 220))
        add_text(slide, chip_x, chip_y + 0.12, chip_w, 0.2, txt, size=8, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    add_text(
        slide,
        0.45,
        0.16,
        12.2,
        0.52,
        "Product Definition: What Is Being Delivered",
        size=30,
        bold=True,
        color=TITLE,
    )
    add_text(
        slide,
        0.45,
        0.74,
        12.2,
        0.45,
        "A standard family detached home with fixed minimum specs for quality consistency and delivery speed.",
        size=12,
        color=SUB,
    )

    draw_floorplan_panel(slide)
    draw_spec_panel(slide)

    add_rect(slide, 0.45, 6.52, 12.4, 0.6, CHARCOAL, CHARCOAL)
    add_text(
        slide,
        0.72,
        6.71,
        12.0,
        0.25,
        "Defined product standards turn policy intent into a buildable, auditable housing unit.",
        size=12,
        bold=True,
        color=WHITE,
    )

    add_text(
        slide,
        0.45,
        7.16,
        12.2,
        0.2,
        "Product settings based on income_title_outline.txt and procurement framing in procurement.txt.",
        size=8,
        color=SUB,
    )

    notes = (
        "This slide answers a simple question: what exactly are we delivering?\n\n"
        "Income Title is not a broad subsidy for whatever the market supplies. It is a defined home product with fixed minimum standards.\n\n"
        "That matters for execution. Standardisation improves procurement comparability, reduces design churn, and helps keep quality and cost inside clear bounds.\n\n"
        "For households, these are normal family homes: three bedrooms, two bathrooms, two living zones, single-storey form, and practical liveability.\n\n"
        "For government and delivery partners, the product is measurable and auditable. That supports faster approvals, cleaner tendering, and reliable scale-up."
    )

    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(OUT)


if __name__ == "__main__":
    main()
