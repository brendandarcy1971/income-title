from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = "proposal/income_title_slide_04_1_glossary.pptx"

BG = RGBColor(246, 248, 250)
TITLE = RGBColor(20, 20, 20)
SUB = RGBColor(70, 70, 70)
CHARCOAL = RGBColor(54, 58, 66)
WHITE = RGBColor(255, 255, 255)
CARD = RGBColor(241, 244, 247)
BORDER = RGBColor(220, 225, 230)


def add_text(slide, left, top, width, height, text, size=14, bold=False, color=TITLE):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_shape(slide, shape_type, left, top, width, height, fill, line=None):
    s = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = fill if line is None else line
    return s


def add_glossary_column(slide, left, top, width, height, heading, items):
    add_shape(slide, 1, left, top, width, height, CARD, BORDER)
    add_text(slide, left + 0.12, top + 0.08, width - 0.24, 0.24, heading, size=11, bold=True, color=SUB)
    y = top + 0.34
    for term, desc in items:
        add_text(slide, left + 0.12, y, width - 0.24, 0.18, f"{term}:", size=9, bold=True)
        y += 0.16
        add_text(slide, left + 0.14, y, width - 0.28, 0.25, desc, size=8, color=TITLE)
        y += 0.28


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    add_text(slide, 0.45, 0.16, 12.3, 0.5, "Slide 4.1: Glossary for Cost and Lending Terms", size=28, bold=True)
    add_text(
        slide,
        0.45,
        0.72,
        12.0,
        0.35,
        "Definitions for the terms used in the Slide 4 waterfall and mortgage-capacity panel.",
        size=12,
        color=SUB,
    )

    left_items = [
        ("Raw land acquisition", "Purchase price of undeveloped land before rezoning and servicing."),
        ("Rezoning/scarcity uplift", "Value uplift from planning status, constrained supply, and permissions."),
        ("Infrastructure and civil works", "Site servicing and subdivision works including roads, utilities, and drainage."),
        ("Construction hard cost", "Direct dwelling build costs: materials, labor, and core subcontract scope."),
        ("Taxes/fees/compliance", "Statutory charges, applications, and regulatory compliance costs."),
        ("Overheads/risk/margin", "Overhead allocation, risk pricing, and commercial margin."),
        ("Final product price", "End market price paid by the buyer for completed home and lot."),
    ]

    right_items = [
        ("Lending multiple", "Simplified borrowing-capacity ratio based on gross income (for example 5x)."),
        ("Maximum borrowing capacity", "Highest indicative loan amount permitted by lender serviceability settings."),
        ("Deposit", "Upfront buyer equity contributed to purchase."),
        ("Loan required", "Purchase price minus deposit."),
        ("Mortgage interest rate", "Annual loan rate used to calculate repayments."),
        ("Loan term", "Repayment duration (for example 30 years)."),
        ("Monthly repayment", "Periodic mortgage payment implied by loan amount, rate, and term."),
        ("Repayment burden", "Annual mortgage repayments divided by gross household income."),
    ]

    add_glossary_column(slide, 0.45, 1.22, 6.25, 5.1, "Cost-Formation Terms", left_items)
    add_glossary_column(slide, 6.95, 1.22, 5.95, 5.1, "Lending and Repayment Terms", right_items)

    add_shape(slide, 1, 0.45, 6.42, 12.45, 0.52, CHARCOAL)
    add_text(
        slide,
        0.70,
        6.57,
        11.9,
        0.24,
        "Slide 4 is an illustrative system-economics view: it explains price formation mechanics and financing pressure points, not a single universal market outcome.",
        size=10,
        bold=True,
        color=WHITE,
    )
    add_text(
        slide,
        0.45,
        7.02,
        12.2,
        0.2,
        "Caveat: calculator outputs are stylized and exclude insurance, rates, maintenance, transaction costs, and lender buffer-policy variations.",
        size=8,
        color=SUB,
    )

    notes = (
        "This glossary slide defines the terms used in Slide 4 so interpretation is consistent.\n\n"
        "On the left, the waterfall terms describe where cost layers arise from raw land through delivery and sale.\n\n"
        "On the right, the lending terms explain how borrowing-capacity assumptions translate into loan size and repayment burden.\n\n"
        "The key point is that Slide 4 combines two systems: cost formation and financing constraints. Understanding both is necessary to understand why final prices can move away from core build cost.\n\n"
        "Treat the numbers as an illustrative scenario, not a universal constant across all submarkets."
    )
    notes_slide = slide.notes_slide
    notes_frame = notes_slide.notes_text_frame
    notes_frame.clear()
    notes_frame.text = notes

    prs.save(OUT)


if __name__ == "__main__":
    main()
